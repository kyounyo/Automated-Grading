import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

df = pd.read_csv('q22_full130_compiled.csv')

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Q22 AI vs Human (130)"

headers = [
    "Response ID", "Human Score", "Gemini Score", "Nemotron Score", "Claude Score",
    "Max Score", "AI Agreement Range", "AI Average", "Gap (Human - AI avg)",
    "Flag", "Student Answer (verbatim)"
]
ws.append(headers)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center")
ws.freeze_panes = "A2"

flag_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
missing_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

row_num = 2
for _, r in df.iterrows():
    rid = int(r['response_id'])
    human = r['human_score']
    gem = r['gemini_score']
    nem = r['nemotron_score']
    claude_missing = pd.isna(r['claude_score'])
    claude_val = None if claude_missing else r['claude_score']
    max_score = r['max_score']
    answer = str(r['student_answer']) if pd.notna(r['student_answer']) else ""

    ws.cell(row=row_num, column=1, value=rid)
    ws.cell(row=row_num, column=2, value=human)
    ws.cell(row=row_num, column=3, value=gem)
    ws.cell(row=row_num, column=4, value=nem)
    ws.cell(row=row_num, column=5, value=claude_val if not claude_missing else "N/A (grading failed)")
    ws.cell(row=row_num, column=6, value=max_score)

    c_col = f"C{row_num}"
    d_col = f"D{row_num}"
    e_col = f"E{row_num}"
    b_col = f"B{row_num}"

    if claude_missing:
        # Only 2 AI scores available -- compute range/avg/gap from Gemini+Nemotron only,
        # and never auto-flag this row (can't claim "3-model consensus" with only 2).
        ws.cell(row=row_num, column=7, value=f"=MAX({c_col},{d_col})-MIN({c_col},{d_col})")
        ws.cell(row=row_num, column=8, value=f"=AVERAGE({c_col},{d_col})")
        ws.cell(row=row_num, column=9, value=f"={b_col}-H{row_num}")
        ws.cell(row=row_num, column=10, value="Claude grading failed -- only 2/3 models available, not auto-flagged")
        ws.cell(row=row_num, column=5).fill = missing_fill
    else:
        ws.cell(row=row_num, column=7, value=f"=MAX({c_col}:{e_col})-MIN({c_col}:{e_col})")
        ws.cell(row=row_num, column=8, value=f"=AVERAGE({c_col}:{e_col})")
        ws.cell(row=row_num, column=9, value=f"={b_col}-H{row_num}")
        ws.cell(row=row_num, column=10,
                value=f'=IF(AND(G{row_num}<=1,I{row_num}>=2),"AI consensus, far below human","")')

    ws.cell(row=row_num, column=11, value=answer)
    row_num += 1

last_row = row_num - 1
# Conditional highlight: color the whole row pink when Flag is non-empty
from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add(
    f"A2:K{last_row}",
    FormulaRule(formula=[f'$J2="AI consensus, far below human"'], fill=flag_fill)
)

col_widths = [12, 11, 11, 13, 11, 10, 15, 11, 18, 32, 70]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ---- Notes sheet ----
notes = wb.create_sheet("Notes")
notes_lines = [
    "Q22 - AI grading (Gemini 3.1 Flash Lite, Nemotron 3 Super 120B, Claude 4.6 Sonnet) vs human marker, FULL 130-student set",
    "",
    "Source of human scores: evaluation/Dataset for prompt.xlsx, 'Response' sheet, Q22 rows (grade column). "
    "Cross-checked against the lecturer's own PHR1012 S2 2023 Grades (Q22) export -- the first ~29 rows matched exactly.",
    "",
    "Source of AI scores: evaluation/results_DYNAMIC_<model>.xlsx, 'Q22_(130_Students)' tab -- the dynamic "
    "Rubric Interpreter architecture, run against the full 130-student Q22 population (the first 25 of these "
    "were the original stratified sample; the remaining 105 were graded in a follow-up run using the same code).",
    "",
    "Claude 4.6 Sonnet: response_id 32885792 failed grading after all retries and has no score -- shown as "
    "'N/A (grading failed)' and excluded from the AI Agreement Range / AI Average / Flag formulas for that row "
    "(computed from Gemini+Nemotron only, and never auto-flagged, since a 2-model comparison cannot support a "
    "'3-model consensus' claim).",
    "",
    "Flag column: marks a response 'AI consensus, far below human' when AI Agreement Range <= 1 mark "
    "(the 3 models broadly agree with each other) AND Gap (Human - AI avg) >= 2 marks (the human score is "
    "substantially higher than what all 3 models independently converged on). These are the cases worth a "
    "second human look -- not because the AI is necessarily right, but because three independent models "
    "reaching similar conclusions against a rubric is a stronger signal than any one model's score alone.",
]
for line in notes_lines:
    notes.append([line])
notes.column_dimensions["A"].width = 120
for row in notes.iter_rows():
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save("Q22_AI_vs_Human_Comparison_Full130.xlsx")
print("Saved Q22_AI_vs_Human_Comparison_Full130.xlsx")

flagged = df.copy()
